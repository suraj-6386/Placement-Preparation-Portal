"""
SkillPrep Portal - REST API Endpoints
Implements authentication, user profiles, aptitude, coding, interview, and help routes.
"""
import os
import re
import json
import html
import io
import secrets
import time
import urllib.parse
import urllib.request
import urllib.error
from datetime import datetime, timedelta
from pathlib import Path
from typing import Optional, Any, List
from fastapi import APIRouter, Depends, Form, File, UploadFile, Header, Query, status
from fastapi.responses import JSONResponse, HTMLResponse, RedirectResponse
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from config import settings
from database import get_db
from models import (
    User,
    SessionModel,
    HelpQuery,
    FAQ,
    InterviewQuestion,
    LearningResource,
    PracticeQuestion,
    MockTest,
    CodingChallenge,
    ActivityEvent,
)
from schemas import (
    UserRegisterRequest,
    UserLoginRequest,
    GoogleAuthRequest,
    GoogleClientIdResponse,
    LogoutRequest,
    AuthResponse,
    BaseResponse,
    UserProfileResponse,
    HelpQueryCreate,
    HelpQueryResponse,
    ActivityCreate,
    InterviewAIRequest,
)
from auth import (
    hash_password,
    verify_password,
    generate_session_token,
    generate_uuid,
    session_is_valid,
    verify_google_token,
    get_current_user,
)

router = APIRouter(prefix="/api", tags=["API"])

BACKEND_DIR = Path(__file__).resolve().parent
DATASET_DIR = BACKEND_DIR / "dataset"
USERDATA_DIR = BACKEND_DIR / "userdata"

# ============================================================================
# File Upload Helpers
# ============================================================================

def secure_filename(filename: str) -> str:
    """Sanitize filename to prevent directory traversal."""
    filename = re.sub(r"[^\w\.-]", "_", os.path.basename(filename))
    return filename.strip("._")


def is_allowed_file(filename: str, allowed_extensions: set) -> bool:
    """Check if uploaded file has an allowed extension."""
    if "." not in filename:
        return False
    ext = filename.rsplit(".", 1)[1].lower()
    return ext in allowed_extensions


async def save_uploaded_photo(user_id: str, file: UploadFile) -> Optional[str]:
    """Save user profile image to userdata/profile_images."""
    if not file or not file.filename:
        return None
    if not is_allowed_file(file.filename, settings.ALLOWED_EXTENSIONS_IMAGES):
        return None

    target_dir = USERDATA_DIR / "profile_images"
    target_dir.mkdir(parents=True, exist_ok=True)
    safe_name = secure_filename(file.filename)
    stored_name = f"{user_id}_{safe_name}"
    target_path = target_dir / stored_name

    content = await file.read()
    with open(target_path, "wb") as f:
        f.write(content)
    return stored_name


async def save_uploaded_resume(user_id: str, file: UploadFile) -> Optional[str]:
    """Save user resume document to userdata/resumes."""
    if not file or not file.filename:
        return None
    if not is_allowed_file(file.filename, settings.ALLOWED_EXTENSIONS_RESUMES):
        return None

    target_dir = USERDATA_DIR / "resumes"
    target_dir.mkdir(parents=True, exist_ok=True)
    safe_name = secure_filename(file.filename)
    stored_name = f"{user_id}_{safe_name}"
    target_path = target_dir / stored_name

    content = await file.read()
    if not content or len(content) > settings.RESUME_MAX_FILE_SIZE:
        return None
    with open(target_path, "wb") as f:
        f.write(content)
    return stored_name


# ============================================================================
# Dataset Fallback Loader
# ============================================================================

def _read_json_file(filename: str, default: Any = None) -> Any:
    """Read a JSON dataset file from the dataset directory."""
    filepath = DATASET_DIR / filename
    if filepath.exists():
        try:
            with open(filepath, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            pass
    return default if default is not None else {}


# ============================================================================
# Authentication Endpoints
# ============================================================================

@router.post("/register", response_model=BaseResponse)
def register(req: UserRegisterRequest, db: Session = Depends(get_db)):
    """
    Register a new user account.
    Validates unique email and username, hashes password with bcrypt,
    and stores user in MySQL.
    """
    if not req.name or not req.email or not req.username or not req.password:
        return JSONResponse(
            status_code=status.HTTP_400_BAD_REQUEST,
            content={"success": False, "message": "All required fields must be provided"},
        )

    username = req.username.strip()
    email = req.email.strip().lower()

    # Check for existing username
    if db.query(User).filter(User.username == username).first():
        return JSONResponse(
            status_code=status.HTTP_400_BAD_REQUEST,
            content={"success": False, "message": "Username already exists"},
        )

    # Check for existing email
    if db.query(User).filter(User.email == email).first():
        return JSONResponse(
            status_code=status.HTTP_400_BAD_REQUEST,
            content={"success": False, "message": "Email already registered"},
        )

    # Create new User model instance
    new_user = User(
        id=generate_uuid(),
        name=req.name.strip(),
        email=email,
        username=username,
        password=hash_password(req.password),
        phone=req.phone or "",
        college=req.college or "",
        course=req.course or "",
        skills=req.skills or "",
        photo="",
        resume="",
    )
    db.add(new_user)
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        return JSONResponse(
            status_code=status.HTTP_400_BAD_REQUEST,
            content={"success": False, "message": "Username or email already exists"},
        )
    db.refresh(new_user)

    return {"success": True, "message": "Registration successful"}


@router.post("/login", response_model=AuthResponse)
def login(req: UserLoginRequest, db: Session = Depends(get_db)):
    """
    Authenticate user with username and password.
    Creates a new session token in MySQL and returns user info.
    """
    if not req.username or not req.password:
        return JSONResponse(
            status_code=status.HTTP_400_BAD_REQUEST,
            content={"success": False, "message": "Username and password required"},
        )

    user = db.query(User).filter(User.username == req.username.strip()).first()
    if not user or not verify_password(req.password, user.password):
        return JSONResponse(
            status_code=status.HTTP_401_UNAUTHORIZED,
            content={"success": False, "message": "Invalid username or password"},
        )

    # Generate session token and store in MySQL sessions table
    token = generate_session_token()
    session = SessionModel(token=token, user_id=user.id, username=user.username)
    db.add(session)
    db.commit()

    return {
        "success": True,
        "message": "Login successful",
        "token": token,
        "username": user.username,
        "name": user.name,
    }


@router.post("/auth/google", response_model=AuthResponse)
def google_auth(req: GoogleAuthRequest, db: Session = Depends(get_db)):
    """
    Authenticate or register user using Google Sign-In (OAuth ID token).
    Verifies the Google credential, finds or creates the user in MySQL,
    and returns an application session token.
    """
    if not req.credential:
        return JSONResponse(
            status_code=status.HTTP_400_BAD_REQUEST,
            content={"success": False, "message": "Google credential token required"},
        )

    try:
        if not settings.GOOGLE_CLIENT_ID:
            return JSONResponse(
                status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
                content={"success": False, "message": "Google Sign-In is not configured"},
            )
        idinfo = verify_google_token(req.credential, settings.GOOGLE_CLIENT_ID or None)
    except Exception as e:
        return JSONResponse(
            status_code=status.HTTP_401_UNAUTHORIZED,
            content={"success": False, "message": f"Google authentication failed: {str(e)}"},
        )

    email = idinfo.get("email")
    if not email:
        return JSONResponse(
            status_code=status.HTTP_400_BAD_REQUEST,
            content={"success": False, "message": "Email not provided by Google account"},
        )

    email = email.lower().strip()
    name = idinfo.get("name") or email.split("@")[0]
    picture = idinfo.get("picture") or ""

    # Check if user already exists
    user = db.query(User).filter(User.email == email).first()

    if not user:
        # Create username based on email
        base_username = email.split("@")[0]
        base_username = re.sub(r"[^\w]", "_", base_username)
        username = base_username

        # Avoid username collision
        counter = 1
        while db.query(User).filter(User.username == username).first():
            username = f"{base_username}_{counter}"
            counter += 1

        # Secure random password hash for OAuth accounts
        random_pwd = generate_session_token()
        user = User(
            id=generate_uuid(),
            name=name,
            email=email,
            username=username,
            password=hash_password(random_pwd),
            phone="",
            college="",
            course="",
            skills="",
            photo=picture,
            resume="",
        )
        db.add(user)
        db.commit()
        db.refresh(user)

    # Issue session token
    token = generate_session_token()
    session = SessionModel(token=token, user_id=user.id, username=user.username)
    db.add(session)
    db.commit()

    return {
        "success": True,
        "message": "Google Sign-In successful",
        "token": token,
        "username": user.username,
        "name": user.name,
    }


@router.get("/auth/google/login")
def google_login_redirect():
    """
    Redirect user to Google's OAuth2 authorization screen.
    Uses configured GOOGLE_CLIENT_ID and GOOGLE_REDIRECT_URI.
    """
    if not settings.GOOGLE_CLIENT_ID:
        return HTMLResponse(
            """
            <html><body style="font-family:sans-serif;padding:3rem;text-align:center;">
            <h2 style="color:#ef4444;">Google Client ID Not Configured</h2>
            <p>Please enter your <code>GOOGLE_CLIENT_ID</code> in <code>.env</code> as described in <code>update.txt</code>.</p>
            <a href="/" style="display:inline-block;padding:8px 16px;background:#4f46e5;color:#fff;text-decoration:none;border-radius:6px;margin-top:1rem;">Back to Home</a>
            </body></html>
            """,
            status_code=400,
        )

    params = {
        "client_id": settings.GOOGLE_CLIENT_ID,
        "redirect_uri": settings.GOOGLE_REDIRECT_URI,
        "response_type": "code",
        "scope": "openid email profile",
        "access_type": "online",
        "prompt": "select_account",
    }
    auth_url = f"https://accounts.google.com/o/oauth2/v2/auth?{urllib.parse.urlencode(params)}"
    return RedirectResponse(url=auth_url)


@router.get("/auth/google/callback", response_class=HTMLResponse)
def google_callback(
    code: Optional[str] = None,
    error: Optional[str] = None,
    db: Session = Depends(get_db),
):
    """
    Handle Google OAuth2 callback.
    Exchanges the authorization code for an access token,
    fetches the user profile from Google, finds or creates the user in MySQL,
    creates a session in MySQL, and sets the auth session in the client.
    """
    if error:
        safe_error = html.escape(error, quote=True)
        return HTMLResponse(
            f"""
            <!DOCTYPE html>
            <html><body style="font-family:sans-serif;padding:3rem;text-align:center;">
            <h2 style="color:#ef4444;">Google Authentication Failed</h2>
            <p style="color:#64748b;">{safe_error}</p>
            <a href="/" style="display:inline-block;padding:8px 18px;background:#4f46e5;color:#fff;text-decoration:none;border-radius:6px;margin-top:1rem;">Return to Portal</a>
            </body></html>
            """,
            status_code=400,
        )

    if not code:
        return HTMLResponse(
            """
            <!DOCTYPE html>
            <html><body style="font-family:sans-serif;padding:3rem;text-align:center;">
            <h2 style="color:#ef4444;">Authorization Code Missing</h2>
            <p style="color:#64748b;">No authorization code received from Google.</p>
            <a href="/" style="display:inline-block;padding:8px 18px;background:#4f46e5;color:#fff;text-decoration:none;border-radius:6px;margin-top:1rem;">Return to Portal</a>
            </body></html>
            """,
            status_code=400,
        )

    try:
        # 1. Exchange authorization code for tokens
        token_url = "https://oauth2.googleapis.com/token"
        token_payload = urllib.parse.urlencode({
            "code": code,
            "client_id": settings.GOOGLE_CLIENT_ID,
            "client_secret": settings.GOOGLE_CLIENT_SECRET,
            "redirect_uri": settings.GOOGLE_REDIRECT_URI,
            "grant_type": "authorization_code",
        }).encode("utf-8")

        req = urllib.request.Request(
            token_url,
            data=token_payload,
            headers={
                "Content-Type": "application/x-www-form-urlencoded",
                "User-Agent": "PlacementPreparationPortal",
            },
        )
        with urllib.request.urlopen(req, timeout=10) as resp:
            token_data = json.loads(resp.read().decode("utf-8"))

        access_token = token_data.get("access_token")
        if not access_token:
            raise ValueError("Failed to obtain access token from Google.")

        # 2. Fetch user information using access token
        userinfo_url = "https://www.googleapis.com/oauth2/v2/userinfo"
        userinfo_req = urllib.request.Request(
            userinfo_url,
            headers={
                "Authorization": f"Bearer {access_token}",
                "User-Agent": "PlacementPreparationPortal",
            },
        )
        with urllib.request.urlopen(userinfo_req, timeout=10) as resp:
            userinfo = json.loads(resp.read().decode("utf-8"))

        email = userinfo.get("email")
        if not email:
            raise ValueError("Google account did not provide an email address.")

        email = email.lower().strip()
        name = userinfo.get("name") or email.split("@")[0]
        picture = userinfo.get("picture") or ""

        # 3. Find or create user in MySQL database
        user = db.query(User).filter(User.email == email).first()
        if not user:
            base_username = re.sub(r"[^\w]", "_", email.split("@")[0])
            username = base_username
            counter = 1
            while db.query(User).filter(User.username == username).first():
                username = f"{base_username}_{counter}"
                counter += 1

            random_pwd = secrets.token_urlsafe(24)
            user = User(
                id=generate_uuid(),
                name=name,
                email=email,
                username=username,
                password=hash_password(random_pwd),
                phone="",
                college="",
                course="",
                skills="",
                photo=picture,
                resume="",
            )
            db.add(user)
            db.commit()
            db.refresh(user)
        elif not user.photo and picture:
            user.photo = picture
            db.commit()

        # 4. Create session token in MySQL sessions table
        session_token = generate_session_token()
        session = SessionModel(token=session_token, user_id=user.id, username=user.username)
        db.add(session)
        db.commit()

        safe_name_html = html.escape(user.name or "", quote=True)
        safe_name_js = json.dumps(user.name or "")
        safe_token_js = json.dumps(session_token)
        return HTMLResponse(
            f"""
            <!DOCTYPE html>
            <html lang="en">
            <head>
                <meta charset="UTF-8">
                <meta name="viewport" content="width=device-width, initial-scale=1.0">
                <title>Authenticating with Google...</title>
                <style>
                    body {{
                        font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif;
                        display: flex;
                        align-items: center;
                        justify-content: center;
                        height: 100vh;
                        margin: 0;
                        background: #f8fafc;
                        color: #1e293b;
                    }}
                    .auth-card {{
                        background: #ffffff;
                        padding: 2.5rem 3rem;
                        border-radius: 16px;
                        box-shadow: 0 10px 25px rgba(0,0,0,0.08);
                        text-align: center;
                        max-width: 420px;
                        width: 90%;
                    }}
                    .spinner {{
                        width: 48px;
                        height: 48px;
                        border: 4px solid #e2e8f0;
                        border-top-color: #4f46e5;
                        border-radius: 50%;
                        animation: spin 0.8s linear infinite;
                        margin: 0 auto 1.5rem;
                    }}
                    @keyframes spin {{
                        to {{ transform: rotate(360deg); }}
                    }}
                </style>
            </head>
            <body>
                <div class="auth-card">
                    <div class="spinner"></div>
                    <h3 style="margin: 0 0 0.5rem; font-weight: 600;">Welcome, {safe_name_html}!</h3>
                    <p style="color: #64748b; margin: 0;">Signed in successfully. Redirecting to portal...</p>
                </div>
                <script>
                    try {{
                        localStorage.setItem('authToken', {safe_token_js});
                        localStorage.setItem('userName', {safe_name_js});
                    }} catch (e) {{
                        console.error('Local storage error', e);
                    }}
                    setTimeout(function() {{
                        window.location.href = '/';
                    }}, 600);
                </script>
            </body>
            </html>
            """
        )
    except Exception as exc:
        return HTMLResponse(
            f"""
            <!DOCTYPE html>
            <html><body style="font-family:sans-serif;padding:3rem;text-align:center;">
            <h2 style="color:#ef4444;">Google Sign-In Error</h2>
            <p style="color:#64748b;">Google Sign-In could not be completed. Please try again.</p>
            <a href="/" style="display:inline-block;padding:8px 18px;background:#4f46e5;color:#fff;text-decoration:none;border-radius:6px;margin-top:1rem;">Return to Portal</a>
            </body></html>
            """,
            status_code=500,
        )


@router.get("/auth/google/client-id", response_model=GoogleClientIdResponse)
def get_google_client_id():
    """
    Retrieve the configured Google OAuth Client ID for frontend initialization.
    """
    return {"client_id": settings.GOOGLE_CLIENT_ID or ""}


@router.post("/logout", response_model=BaseResponse)
def logout(req: LogoutRequest, db: Session = Depends(get_db)):
    """
    Invalidate and remove session token from MySQL.
    """
    if req.token:
        session = db.query(SessionModel).filter(SessionModel.token == req.token).first()
        if session:
            db.delete(session)
            db.commit()
    return {"success": True, "message": "Logout successful"}


# ============================================================================
# User Profile Endpoints
# ============================================================================

def _resolve_token(token: Optional[str], authorization: Optional[str]) -> Optional[str]:
    if token:
        return token
    if authorization:
        parts = authorization.split()
        return parts[1] if len(parts) == 2 and parts[0].lower() == "bearer" else authorization
    return None


def _resume_error(message: str, code: int = status.HTTP_400_BAD_REQUEST):
    return JSONResponse(status_code=code, content={"success": False, "message": message})


def _authenticated_user(token: Optional[str], authorization: Optional[str], db: Session):
    auth_token = _resolve_token(token, authorization)
    if not auth_token:
        return None, _resume_error("Authentication token required", status.HTTP_401_UNAUTHORIZED)
    session = db.query(SessionModel).filter(SessionModel.token == auth_token).first()
    if not session or not session_is_valid(session):
        if session:
            db.delete(session)
            db.commit()
        return None, _resume_error("Invalid or expired session", status.HTTP_401_UNAUTHORIZED)
    user = session.user or db.query(User).filter(User.username == session.username).first()
    if not user:
        return None, _resume_error("User not found", status.HTTP_404_NOT_FOUND)
    return user, None


def _resume_extension(filename: str) -> str:
    return filename.rsplit(".", 1)[1].lower() if "." in filename else ""


async def _read_resume(file: UploadFile, allowed: set) -> tuple[str, bytes]:
    filename = secure_filename(file.filename or "")
    extension = _resume_extension(filename)
    if extension not in allowed:
        raise ValueError("Upload a PDF, DOC, or DOCX file with a valid extension.")
    content = await file.read()
    if not content:
        raise ValueError("The uploaded resume is empty.")
    if len(content) > settings.RESUME_MAX_FILE_SIZE:
        raise ValueError(f"Resume files must be smaller than {settings.RESUME_MAX_FILE_SIZE // (1024 * 1024)} MB.")
    return extension, content


def _docx_to_text(content: bytes) -> str:
    from docx import Document

    document = Document(io.BytesIO(content))
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            paragraphs.extend(
                cell.text.strip()
                for cell in row.cells
                if cell.text.strip()
            )
    text = "\n".join(paragraphs).strip()
    if not text:
        raise ValueError("The uploaded resume does not contain readable text.")
    return text


def _pdf_to_text(content: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(content))
    text = "\n".join((page.extract_text() or "") for page in reader.pages).strip()
    if not text:
        raise ValueError("The PDF does not contain readable text. Scanned PDFs are not supported yet.")
    return text


def _legacy_doc_to_text(content: bytes) -> str:
    """Extract readable text from common binary .doc files without exposing them to a converter."""
    candidates = [content.decode("utf-16le", errors="ignore"), content.decode("cp1252", errors="ignore")]
    fragments = []
    for candidate in candidates:
        fragments.extend(re.findall(r"[A-Za-z0-9][A-Za-z0-9 .,:;()/+&@#%_'\-]{3,}", candidate))
    text = "\n".join(dict.fromkeys(fragment.strip() for fragment in fragments if fragment.strip()))
    if len(re.sub(r"[^A-Za-z]", "", text)) < 20:
        raise ValueError("The DOC file does not contain readable text. Save it as DOCX and try again.")
    return text[: settings.RESUME_MAX_TEXT_LENGTH]


def _call_gemini(resume_text: str) -> dict:
    if not settings.GEMINI_API_KEY:
        raise RuntimeError("Resume AI is not configured. Add GEMINI_API_KEY to the server environment.")
    prompt = """Review the following resume for a job seeker. Return only valid JSON with these keys:
overall_score (integer 0-100), ats_readability_score (integer 0-100), strengths (array of strings),
weaknesses (array of strings), missing_sections (array of strings), suggestions (array of strings).
Be specific, practical, and do not invent facts. Resume:\n""" + resume_text[: settings.RESUME_MAX_TEXT_LENGTH]
    payload = json.dumps({
        "contents": [{"parts": [{"text": prompt}]}],
        "generationConfig": {"temperature": 0.2, "responseMimeType": "application/json"},
    }).encode("utf-8")
    url = (
        f"https://generativelanguage.googleapis.com/v1beta/models/{urllib.parse.quote(settings.GEMINI_MODEL)}:generateContent"
        f"?key={urllib.parse.quote(settings.GEMINI_API_KEY)}"
    )
    request = urllib.request.Request(url, data=payload, headers={"Content-Type": "application/json"}, method="POST")
    try:
        with urllib.request.urlopen(request, timeout=30) as response:
            response_data = json.loads(response.read().decode("utf-8"))
    except Exception as exc:
        raise RuntimeError("Resume AI is temporarily unavailable. Please try again.") from exc
    try:
        raw = response_data["candidates"][0]["content"]["parts"][0]["text"]
        result = json.loads(raw)
        for key in ("strengths", "weaknesses", "missing_sections", "suggestions"):
            if not isinstance(result.get(key), list):
                result[key] = []
        result["overall_score"] = max(0, min(100, int(result.get("overall_score", 0))))
        result["ats_readability_score"] = max(0, min(100, int(result.get("ats_readability_score", 0))))
        return result
    except (KeyError, TypeError, ValueError, json.JSONDecodeError) as exc:
        raise RuntimeError("Resume AI returned an invalid review. Please try again.") from exc


def _call_gemini_json(prompt: str) -> dict:
    """Call Gemini for a JSON response used by interview tools."""
    if not settings.GEMINI_API_KEY:
        raise RuntimeError("Interview AI is not configured. Add GEMINI_API_KEY to the server environment.")
    payload = json.dumps({
        "contents": [{"parts": [{"text": prompt}]}],
        "generationConfig": {"temperature": 0.35, "responseMimeType": "application/json"},
    }).encode("utf-8")
    url = (
        f"https://generativelanguage.googleapis.com/v1beta/models/{urllib.parse.quote(settings.GEMINI_MODEL)}:generateContent"
        f"?key={urllib.parse.quote(settings.GEMINI_API_KEY)}"
    )
    request = urllib.request.Request(url, data=payload, headers={"Content-Type": "application/json"}, method="POST")
    try:
        with urllib.request.urlopen(request, timeout=45) as response:
            response_data = json.loads(response.read().decode("utf-8"))
        raw = response_data["candidates"][0]["content"]["parts"][0]["text"].strip()
        if raw.startswith("```"):
            raw = re.sub(r"^```(?:json)?\s*|\s*```$", "", raw, flags=re.IGNORECASE)
        return json.loads(raw)
    except urllib.error.HTTPError as exc:
        if exc.code == 429:
            raise RuntimeError("Interview AI rate limit reached. Please wait a moment and try again.") from exc
        raise RuntimeError("Interview AI is temporarily unavailable. Please try again.") from exc
    except (urllib.error.URLError, TimeoutError, KeyError, TypeError, json.JSONDecodeError) as exc:
        raise RuntimeError("Interview AI returned an invalid response. Please try again.") from exc


ACTIVITY_EVENT_TYPES = {
    "practice_answered",
    "mock_completed",
    "challenge_attempt",
    "challenge_completed",
    "interview_viewed",
    "resource_viewed",
    "resume_reviewed",
}
ACTIVITY_CATEGORIES = {
    "aptitude",
    "coding",
    "aptitude_mock",
    "coding_mock",
    "coding_challenge",
    "interview",
    "learning",
    "resume",
}


def _add_activity_event(
    db: Session,
    user_id: str,
    event_type: str,
    category: str,
    item_key: str,
    score: Optional[int] = None,
    details: Optional[dict] = None,
):
    db.add(ActivityEvent(
        id=generate_uuid(),
        user_id=user_id,
        event_type=event_type,
        category=category,
        item_key=item_key,
        score=score,
        details=details or {},
    ))


@router.post("/activity", response_model=BaseResponse)
def record_activity(
    req: ActivityCreate,
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """Persist a bounded, user-owned learning event from an existing activity surface."""
    event_type = req.event_type.strip().lower()
    category = req.category.strip().lower()
    item_key = req.item_key.strip()
    if event_type not in ACTIVITY_EVENT_TYPES or category not in ACTIVITY_CATEGORIES:
        return JSONResponse(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            content={"success": False, "message": "Unsupported activity event"},
        )
    if not item_key or len(item_key) > 150:
        return JSONResponse(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            content={"success": False, "message": "Activity item is invalid"},
        )
    score = req.score
    if score is not None and not 0 <= score <= 100:
        return JSONResponse(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            content={"success": False, "message": "Activity score must be between 0 and 100"},
        )
    details = req.details if isinstance(req.details, dict) else {}
    _add_activity_event(db, user.id, event_type, category, item_key, score, details)
    db.commit()
    return {"success": True, "message": "Activity recorded"}


def _content_total(db: Session, model, category: str, dataset_name: str) -> int:
    total = db.query(model).filter(model.category == category).count()
    if total:
        return total
    data = _read_json_file(dataset_name, {})
    return sum(len(items) for items in data.values()) if isinstance(data, dict) else 0


def _challenge_content_total(db: Session) -> int:
    database_total = db.query(CodingChallenge).count()
    if database_total:
        return database_total
    data = _read_json_file("coding_challenges.json", {})
    return sum(
        len(items)
        for language in data.values()
        if isinstance(language, dict)
        for key in ("easy", "moderate", "hard")
        for items in [language.get(key, [])]
    )


def _interview_subject_total(db: Session) -> int:
    database_total = db.query(InterviewQuestion.subject).filter(
        InterviewQuestion.category == "technical",
        InterviewQuestion.subject.isnot(None),
    ).distinct().count()
    if database_total:
        return database_total
    data = _read_json_file("interview_technical.json", {})
    return len(data) if isinstance(data, dict) else 0


def _percent(completed: int, total: int) -> int:
    return round(min(1, completed / total) * 100) if total else 0


def _event_label(event: ActivityEvent) -> str:
    labels = {
        "practice_answered": "Answered practice questions",
        "mock_completed": "Completed a mock test",
        "challenge_attempt": "Attempted a coding challenge",
        "challenge_completed": "Completed a coding challenge",
        "interview_viewed": "Explored an interview topic",
        "resource_viewed": "Opened a learning resource",
        "resume_reviewed": "Reviewed your resume with AI",
    }
    return labels.get(event.event_type, "Made progress")


def _dashboard_events(db: Session, user_id: str) -> list[ActivityEvent]:
    return (
        db.query(ActivityEvent)
        .filter(ActivityEvent.user_id == user_id)
        .order_by(ActivityEvent.created_at.desc())
        .all()
    )


@router.get("/dashboard")
def get_dashboard(
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    """Return authenticated progress metrics derived from persisted user activity."""
    events = _dashboard_events(db, user.id)
    practice_latest = {}
    interview_topics = set()
    completed_challenges = set()
    for event in events:
        if event.event_type == "practice_answered":
            practice_latest.setdefault((event.category, event.item_key), event)
        elif event.event_type == "interview_viewed":
            interview_topics.add(event.item_key)
        elif event.event_type == "challenge_completed":
            completed_challenges.add(event.item_key)

    aptitude_answered = sum(1 for key in practice_latest if key[0] == "aptitude")
    coding_answered = sum(1 for key in practice_latest if key[0] == "coding")
    correct_answers = sum(
        1 for event in practice_latest.values() if event.score == 100
    )
    aptitude_total = _content_total(db, PracticeQuestion, "aptitude", "aptitude_practice.json")
    coding_total = _content_total(db, PracticeQuestion, "coding", "coding_practice.json")
    challenge_total = _challenge_content_total(db)
    interview_total = _interview_subject_total(db)
    mock_tests = sum(1 for event in events if event.event_type == "mock_completed")
    challenge_attempts = sum(1 for event in events if event.event_type == "challenge_attempt")
    resume_reviews = sum(1 for event in events if event.event_type == "resume_reviewed")

    profile_fields = (user.college, user.course, user.skills, user.resume, user.photo)
    profile_progress = _percent(sum(bool(value and value.strip()) for value in profile_fields), 5)
    progress = [
        {"key": "profile", "label": "Profile setup", "value": profile_progress, "completed": sum(bool(value and value.strip()) for value in profile_fields), "total": 5, "href": "profile.html"},
        {"key": "aptitude", "label": "Aptitude practice", "value": _percent(aptitude_answered, aptitude_total), "completed": aptitude_answered, "total": aptitude_total, "href": "aptitude.html"},
        {"key": "coding", "label": "Coding practice", "value": _percent(coding_answered, coding_total), "completed": coding_answered, "total": coding_total, "href": "coding.html"},
        {"key": "challenges", "label": "Coding challenges", "value": _percent(len(completed_challenges), challenge_total), "completed": len(completed_challenges), "total": challenge_total, "href": "coding.html"},
        {"key": "interview", "label": "Interview topics", "value": _percent(len(interview_topics), interview_total), "completed": len(interview_topics), "total": interview_total, "href": "interview.html"},
        {"key": "resume", "label": "Resume readiness", "value": 100 if resume_reviews else 50 if user.resume else 0, "completed": 2 if resume_reviews else 1 if user.resume else 0, "total": 2, "href": "resume.html"},
    ]
    weights = {"profile": 15, "aptitude": 20, "coding": 25, "challenges": 15, "interview": 10, "resume": 15}
    overall_progress = round(sum(item["value"] * weights[item["key"]] for item in progress) / 100)

    today = datetime.utcnow().date()
    daily_activity = []
    for offset in range(6, -1, -1):
        day = today - timedelta(days=offset)
        day_events = [event for event in events if event.created_at and event.created_at.date() == day]
        daily_activity.append({"date": day.isoformat(), "label": day.strftime("%a"), "count": len(day_events)})

    active_days = {event.created_at.date() for event in events if event.created_at}
    streak = 0
    cursor = today
    while cursor in active_days:
        streak += 1
        cursor -= timedelta(days=1)

    recent = [
        {"label": _event_label(event), "item_key": event.item_key, "category": event.category, "score": event.score, "created_at": event.created_at.isoformat() if event.created_at else None}
        for event in events[:8]
    ]
    achievements = [
        {"key": "first-step", "label": "First step", "description": "Complete your first learning activity", "earned": bool(events)},
        {"key": "aptitude-starter", "label": "Aptitude starter", "description": "Answer 10 aptitude questions", "earned": aptitude_answered >= 10},
        {"key": "coding-explorer", "label": "Coding explorer", "description": "Answer 10 coding questions", "earned": coding_answered >= 10},
        {"key": "challenge-solver", "label": "Challenge solver", "description": "Complete a coding challenge", "earned": bool(completed_challenges)},
        {"key": "resume-ready", "label": "Resume ready", "description": "Review your resume with AI", "earned": bool(resume_reviews)},
    ]
    recommendations = []
    if not user.resume:
        recommendations.append({"title": "Upload your resume", "description": "Get a baseline and unlock resume readiness.", "href": "resume.html"})
    if profile_progress < 100:
        recommendations.append({"title": "Finish your profile", "description": "Add your course, skills, and college details.", "href": "profile.html"})
    if aptitude_answered < 10:
        recommendations.append({"title": "Build an aptitude streak", "description": "Answer 10 questions to establish your baseline.", "href": "aptitude.html"})
    if coding_answered < 10:
        recommendations.append({"title": "Practice a coding topic", "description": "Choose a language and solve a few questions.", "href": "coding.html"})
    if len(interview_topics) < interview_total:
        recommendations.append({"title": "Explore interview topics", "description": "Open a technical subject and review its questions.", "href": "interview.html"})

    return {
        "success": True,
        "user": {"name": user.name or "Learner", "resume": bool(user.resume)},
        "overall_progress": overall_progress,
        "kpis": {"questions_answered": aptitude_answered + coding_answered, "correct_answers": correct_answers, "mock_tests": mock_tests, "challenge_attempts": challenge_attempts, "challenges_completed": len(completed_challenges), "resume_reviews": resume_reviews, "active_streak": streak, "activity_count": len(events)},
        "progress": progress,
        "daily_activity": daily_activity,
        "recent_activity": recent,
        "achievements": achievements,
        "recommendations": recommendations[:4],
        "is_new_user": not events,
    }


@router.post("/resume/analyze")
async def analyze_resume(
    file: UploadFile = File(...),
    token: Optional[str] = Query(None),
    authorization: Optional[str] = Header(None),
    db: Session = Depends(get_db),
):
    user, error = _authenticated_user(token, authorization, db)
    if error:
        return error
    try:
        extension, content = await _read_resume(file, settings.ALLOWED_EXTENSIONS_RESUMES)
        if extension == "pdf":
            text = _pdf_to_text(content)
        elif extension == "doc":
            text = _legacy_doc_to_text(content)
        else:
            text = _docx_to_text(content)
        review = _call_gemini(text)
        _add_activity_event(db, user.id, "resume_reviewed", "resume", file.filename or "resume", 100, {
            "overall_score": review.get("overall_score", 0),
        })
        db.commit()
        return {"success": True, "filename": file.filename, "review": review}
    except ValueError as exc:
        return _resume_error(str(exc), status.HTTP_422_UNPROCESSABLE_ENTITY)
    except RuntimeError as exc:
        return _resume_error(str(exc), status.HTTP_503_SERVICE_UNAVAILABLE)
    except Exception:
        return _resume_error("The resume could not be processed.", status.HTTP_422_UNPROCESSABLE_ENTITY)


@router.get("/profile", response_model=UserProfileResponse)
def get_profile(
    token: Optional[str] = Query(None),
    authorization: Optional[str] = Header(None),
    db: Session = Depends(get_db),
):
    """
    Retrieve user profile information by session token.
    """
    auth_token = _resolve_token(token, authorization)
    if not auth_token:
        return JSONResponse(
            status_code=status.HTTP_400_BAD_REQUEST,
            content={"success": False, "message": "Authentication token required"},
        )

    session = db.query(SessionModel).filter(SessionModel.token == auth_token).first()
    if not session or not session_is_valid(session):
        if session:
            db.delete(session)
            db.commit()
        return JSONResponse(
            status_code=status.HTTP_401_UNAUTHORIZED,
            content={"success": False, "message": "Invalid or expired session"},
        )

    user = session.user or db.query(User).filter(User.username == session.username).first()
    if not user:
        return JSONResponse(
            status_code=status.HTTP_404_NOT_FOUND,
            content={"success": False, "message": "User not found"},
        )

    return {"success": True, "user": user.to_dict()}


@router.post("/profile/update", response_model=BaseResponse)
async def update_profile(
    token: Optional[str] = Form(None),
    name: Optional[str] = Form(None),
    email: Optional[str] = Form(None),
    phone: Optional[str] = Form(None),
    college: Optional[str] = Form(None),
    course: Optional[str] = Form(None),
    skills: Optional[str] = Form(None),
    photo: Optional[UploadFile] = File(None),
    resume: Optional[UploadFile] = File(None),
    authorization: Optional[str] = Header(None),
    db: Session = Depends(get_db),
):
    """
    Update user personal details, profile image, and resume in MySQL.
    """
    auth_token = _resolve_token(token, authorization)
    if not auth_token:
        return JSONResponse(
            status_code=status.HTTP_400_BAD_REQUEST,
            content={"success": False, "message": "Authentication token required"},
        )

    session = db.query(SessionModel).filter(SessionModel.token == auth_token).first()
    if not session or not session_is_valid(session):
        if session:
            db.delete(session)
            db.commit()
        return JSONResponse(
            status_code=status.HTTP_401_UNAUTHORIZED,
            content={"success": False, "message": "Invalid or expired session"},
        )

    user = session.user or db.query(User).filter(User.username == session.username).first()
    if not user:
        return JSONResponse(
            status_code=status.HTTP_404_NOT_FOUND,
            content={"success": False, "message": "User not found"},
        )

    # Save uploaded files if provided
    if photo and photo.filename:
        photo_filename = await save_uploaded_photo(user.id, photo)
        if photo_filename:
            user.photo = photo_filename

    if resume and resume.filename:
        resume_filename = await save_uploaded_resume(user.id, resume)
        if resume_filename:
            user.resume = resume_filename

    # Update personal details
    if name is not None:
        user.name = name.strip()
    if email is not None:
        normalized_email = email.strip().lower()
        existing_email = (
            db.query(User)
            .filter(User.email == normalized_email, User.id != user.id)
            .first()
        )
        if existing_email:
            return JSONResponse(
                status_code=status.HTTP_400_BAD_REQUEST,
                content={"success": False, "message": "Email already registered"},
            )
        user.email = normalized_email
    if phone is not None:
        user.phone = phone.strip()
    if college is not None:
        user.college = college.strip()
    if course is not None:
        user.course = course.strip()
    if skills is not None:
        user.skills = skills.strip()

    db.commit()
    db.refresh(user)

    return {"success": True, "message": "Profile updated successfully"}


# ============================================================================
# Aptitude Endpoints
# ============================================================================

@router.get("/aptitude/learn")
def get_aptitude_learn(db: Session = Depends(get_db)):
    """Curated learning resources and tutorial links for aptitude."""
    items = db.query(LearningResource).filter(LearningResource.category == "aptitude").all()
    if items:
        return {item.title: item.url for item in items}
    return _read_json_file("aptitude_learn.json", {})


@router.get("/aptitude/practice")
def get_aptitude_practice(db: Session = Depends(get_db)):
    """Chapter-wise aptitude practice questions."""
    items = db.query(PracticeQuestion).filter(PracticeQuestion.category == "aptitude").all()
    if items:
        result = {}
        for item in items:
            result.setdefault(item.chapter, []).append(item.to_dict())
        return result
    return _read_json_file("aptitude_practice.json", {})


@router.get("/aptitude/mock")
def get_aptitude_mock(db: Session = Depends(get_db)):
    """Aptitude mock test sets categorized by difficulty level."""
    items = (
        db.query(MockTest)
        .filter(MockTest.category == "aptitude")
        .order_by(MockTest.set_index.asc())
        .all()
    )
    if items:
        result = {}
        for item in items:
            result.setdefault(item.level, {}).setdefault("sets", []).append(item.to_dict())
        return result
    return _read_json_file("aptitude_mock.json", {})


# ============================================================================
# Coding Endpoints
# ============================================================================

@router.get("/coding/learn")
def get_coding_learn(db: Session = Depends(get_db)):
    """Programming languages learning resources and documentation."""
    items = db.query(LearningResource).filter(LearningResource.category == "coding").all()
    if items:
        return {item.title: item.url for item in items}
    return _read_json_file("coding_learn.json", {})


@router.get("/coding/practice")
def get_coding_practice(db: Session = Depends(get_db)):
    """Language-wise coding practice questions."""
    items = db.query(PracticeQuestion).filter(PracticeQuestion.category == "coding").all()
    if items:
        result = {}
        for item in items:
            result.setdefault(item.chapter, []).append(item.to_dict())
        return result
    return _read_json_file("coding_practice.json", {})


@router.get("/coding/mock")
def get_coding_mock(db: Session = Depends(get_db)):
    """Coding mock test sets categorized by difficulty."""
    items = (
        db.query(MockTest)
        .filter(MockTest.category == "coding")
        .order_by(MockTest.set_index.asc())
        .all()
    )
    if items:
        result = {}
        for item in items:
            result.setdefault(item.level, {}).setdefault("sets", []).append(item.to_dict())
        return result
    return _read_json_file("coding_mock.json", {})


LANGUAGE_METADATA = {
    "Python": {"icon": "🐍", "color": "#3776AB"},
    "Java": {"icon": "☕", "color": "#007396"},
    "C": {"icon": "©️", "color": "#A8B9CC"},
    "C++": {"icon": "⚡", "color": "#00599C"},
    "JavaScript": {"icon": "📜", "color": "#F7DF1E"},
    "PHP": {"icon": "🐘", "color": "#777BB4"},
    "CSharp": {"icon": "#️⃣", "color": "#239120", "displayName": "C#"},
    "Frontend": {"icon": "🎨", "color": "#FF6B6B"},
    "Backend": {"icon": "⚙️", "color": "#4ECDC4"},
    "SQL": {"icon": "🗄️", "color": "#336791"},
}


@router.get("/coding/challenges")
def get_coding_challenges(db: Session = Depends(get_db)):
    """Coding challenges across 10 programming languages."""
    items = db.query(CodingChallenge).all()
    if items:
        result = {}
        for item in items:
            lang = item.language
            if lang not in result:
                meta = LANGUAGE_METADATA.get(lang, {"icon": "💻", "color": "#3b82f6"})
                result[lang] = {
                    "icon": meta.get("icon", "💻"),
                    "color": meta.get("color", "#3b82f6"),
                    "easy": [],
                    "moderate": [],
                    "hard": [],
                }
                if "displayName" in meta:
                    result[lang]["displayName"] = meta["displayName"]
            diff = (item.difficulty or "easy").lower()
            if diff not in ("easy", "moderate", "hard"):
                diff = "easy"
            result[lang][diff].append(item.to_dict())
        return result
    return _read_json_file("coding_challenges.json", {})


# ============================================================================
# Interview Endpoints
# ============================================================================

INTERVIEW_DIFFICULTIES = ("Basic", "Medium", "Hard")
INTERVIEW_AI_LIMIT = 20
_INTERVIEW_AI_REQUESTS: dict[str, list[float]] = {}
HR_EXTRA_QUESTIONS = [
    "How would your manager describe you?",
    "What type of feedback helps you improve?",
    "Tell me about a time you disagreed with a teammate.",
    "How do you manage competing deadlines?",
    "Describe a time you had to learn something quickly.",
    "How do you build relationships with new teammates?",
    "Tell me about a time you went beyond expectations.",
    "What kind of manager helps you do your best work?",
    "How do you respond when priorities suddenly change?",
    "Describe a time you had to persuade someone.",
    "How do you make decisions with incomplete information?",
    "Tell me about a time you solved a problem creatively.",
    "What does success mean to you in this role?",
    "How do you make sure your work is accurate?",
    "Describe a time you helped a struggling teammate.",
    "What would you do in your first 30 days here?",
    "How do you handle a task you do not enjoy?",
    "Tell me about a time you improved a process.",
    "How do you communicate bad news professionally?",
    "What are you looking for in your next opportunity?",
    "How do you balance speed and quality?",
    "Describe a time you had to work with limited resources.",
    "What is one skill you are actively developing?",
    "How do you prepare before an important presentation?",
    "Why should we choose you over another candidate?",
]


def _interview_question_data(db: Session) -> dict:
    items = (
        db.query(InterviewQuestion)
        .filter(InterviewQuestion.category == "technical", InterviewQuestion.subject != "Go")
        .order_by(InterviewQuestion.display_order.asc(), InterviewQuestion.id.asc())
        .all()
    )
    result = {}
    for item in items:
        result.setdefault(item.subject or "General", []).append(item.question)
    if not result:
        result = _read_json_file("interview_technical.json", {})
    result.pop("Go", None)
    result.setdefault("API", [
        "What is an API and how does a client consume one?",
        "Compare REST, GraphQL, and SOAP APIs.",
        "What are HTTP methods and when should each be used?",
        "How do authentication and authorization differ in an API?",
        "How would you version a public API?",
        "What makes an API idempotent?",
        "How do you design useful API error responses?",
        "What is rate limiting and why is it important?",
        "How do you secure sensitive data in API requests?",
        "How would you test an API in CI?",
    ])
    return result


def _check_interview_ai_rate_limit(user_id: str):
    now = time.monotonic()
    recent = [stamp for stamp in _INTERVIEW_AI_REQUESTS.get(user_id, []) if now - stamp < 600]
    if len(recent) >= INTERVIEW_AI_LIMIT:
        raise RuntimeError("Interview AI rate limit reached. Please wait a few minutes and try again.")
    recent.append(now)
    _INTERVIEW_AI_REQUESTS[user_id] = recent


def _fallback_interview_questions(category: str, difficulty: str, seeds: list[str]) -> list[str]:
    questions = list(dict.fromkeys(seeds))
    prompts = {
        "Basic": ["Define the core concept and give a simple example.", "What are the main benefits and limitations?", "How would you explain this to a beginner?"],
        "Medium": ["Compare two practical approaches and explain the trade-offs.", "How would you debug or test this in a real project?", "Describe a realistic implementation scenario."],
        "Hard": ["Design a robust solution for a high-scale production scenario.", "Analyze edge cases, failure modes, and performance trade-offs.", "How would you improve this design under strict reliability requirements?"],
    }
    index = 0
    while len(questions) < 50:
        seed = seeds[index % len(seeds)] if seeds else f"{category} interview concept"
        prompt = prompts[difficulty][index % len(prompts[difficulty])]
        questions.append(f"{seed} {prompt}")
        index += 1
    return questions[:50]


def _get_interview_question_bank(category: str, difficulty: str, db: Session) -> tuple[list[str], str]:
    catalog = _interview_question_data(db)
    if category not in catalog:
        raise ValueError("Unknown technical interview category")
    if difficulty not in INTERVIEW_DIFFICULTIES:
        raise ValueError("Difficulty must be Basic, Medium, or Hard")
    cache_dir = USERDATA_DIR / "interview_question_cache"
    cache_dir.mkdir(parents=True, exist_ok=True)
    cache_path = cache_dir / f"{secure_filename(category)}_{difficulty.lower()}.json"
    if cache_path.exists():
        try:
            cached = json.loads(cache_path.read_text(encoding="utf-8"))
            if isinstance(cached, list) and len(cached) == 50:
                return cached, "ai-cache"
        except (OSError, ValueError):
            pass
    seeds = [str(item) for item in catalog[category] if str(item).strip()]
    prompt = f"""Generate exactly 50 unique {difficulty}-level technical interview questions for {category}.
Return only JSON in this shape: {{\"questions\":[\"question 1\", ...]}}.
""" + "\nSeed topics:\n" + "\n".join(f"- {seed}" for seed in seeds[:30])
    source = "ai"
    try:
        generated = _call_gemini_json(prompt)
        questions = generated.get("questions", []) if isinstance(generated, dict) else []
        questions = [str(question).strip() for question in questions if str(question).strip()]
        if len(questions) < 50:
            raise RuntimeError("The AI question bank was incomplete.")
        questions = list(dict.fromkeys(questions))[:50]
        if len(questions) < 50:
            raise RuntimeError("The AI question bank contained duplicates.")
    except RuntimeError:
        questions = _fallback_interview_questions(category, difficulty, seeds)
        source = "fallback"
    try:
        cache_path.write_text(json.dumps(questions, ensure_ascii=False, indent=2), encoding="utf-8")
    except OSError:
        pass
    return questions, source

@router.get("/interview/hr")
def get_interview_hr(db: Session = Depends(get_db)):
    """Common HR interview questions and behavioral prompts."""
    items = (
        db.query(InterviewQuestion)
        .filter(InterviewQuestion.category == "hr")
        .order_by(InterviewQuestion.display_order.asc(), InterviewQuestion.id.asc())
        .all()
    )
    questions = [item.question for item in items] if items else _read_json_file("interview_hr.json", [])
    questions = list(dict.fromkeys(questions + HR_EXTRA_QUESTIONS))
    return questions[:50]


@router.get("/interview/technical")
def get_interview_technical(db: Session = Depends(get_db)):
    """Subject-wise technical interview seed questions, excluding Go and including API."""
    return _interview_question_data(db)


@router.get("/interview/catalog")
def get_interview_catalog(db: Session = Depends(get_db)):
    return {"categories": sorted(_interview_question_data(db).keys()), "difficulties": list(INTERVIEW_DIFFICULTIES)}


@router.get("/interview/questions")
def get_interview_questions(
    category: str = Query(..., min_length=1, max_length=80),
    difficulty: str = Query(..., min_length=1, max_length=20),
    db: Session = Depends(get_db),
):
    try:
        questions, source = _get_interview_question_bank(category.strip(), difficulty.strip().title(), db)
        return {"success": True, "category": category, "difficulty": difficulty, "source": source, "questions": questions}
    except ValueError as exc:
        return _resume_error(str(exc), status.HTTP_422_UNPROCESSABLE_ENTITY)
    except RuntimeError as exc:
        return _resume_error(str(exc), status.HTTP_503_SERVICE_UNAVAILABLE)


@router.post("/interview/evaluate")
def evaluate_interview_answer(req: InterviewAIRequest, user: User = Depends(get_current_user)):
    question = req.question.strip()
    answer = (req.answer or "").strip()
    if not question or len(question) > 1000 or not answer:
        return _resume_error("Provide a question and a non-empty answer.", status.HTTP_422_UNPROCESSABLE_ENTITY)
    if len(answer) > 8000:
        return _resume_error("Your answer is too long. Keep it under 8,000 characters.", status.HTTP_422_UNPROCESSABLE_ENTITY)
    try:
        _check_interview_ai_rate_limit(user.id)
        result = _call_gemini_json(f"""Evaluate this {req.category} interview answer for the question below.
Return only JSON with integer score (0-100), string feedback, array strengths, and array improvements.
Be constructive, specific, and judge relevance, structure, clarity, evidence, and professionalism.
Question: {question}
Answer: {answer}""")
        score = max(0, min(100, int(result.get("score", 0))))
        return {"success": True, "score": score, "feedback": str(result.get("feedback", "")), "strengths": result.get("strengths", []), "improvements": result.get("improvements", [])}
    except RuntimeError as exc:
        code = status.HTTP_429_TOO_MANY_REQUESTS if "rate limit" in str(exc).lower() else status.HTTP_503_SERVICE_UNAVAILABLE
        return _resume_error(str(exc), code)


@router.post("/interview/best-answer")
def get_best_interview_answer(req: InterviewAIRequest, user: User = Depends(get_current_user)):
    question = req.question.strip()
    if not question or len(question) > 1000:
        return _resume_error("A valid interview question is required.", status.HTTP_422_UNPROCESSABLE_ENTITY)
    try:
        _check_interview_ai_rate_limit(user.id)
        result = _call_gemini_json(f"""Write a strong, natural, interview-appropriate answer to this {req.category} question.
Return only JSON with string answer and array tips. Do not invent personal facts; use clearly marked placeholders when needed.
Question: {question}
Difficulty: {req.difficulty or 'general'}""")
        return {"success": True, "answer": str(result.get("answer", "")), "tips": result.get("tips", [])}
    except RuntimeError as exc:
        code = status.HTTP_429_TOO_MANY_REQUESTS if "rate limit" in str(exc).lower() else status.HTTP_503_SERVICE_UNAVAILABLE
        return _resume_error(str(exc), code)


# ============================================================================
# Help & Support Endpoints
# ============================================================================

@router.get("/faqs")
def get_faqs(db: Session = Depends(get_db)):
    """Frequently Asked Questions."""
    items = db.query(FAQ).all()
    if items:
        return [item.to_dict() for item in items]
    return _read_json_file("faqs.json", [])


@router.post("/help/contact", response_model=HelpQueryResponse)
def submit_help_contact(req: HelpQueryCreate, db: Session = Depends(get_db)):
    """Submit a contact or support query to MySQL."""
    if not req.name or not req.email or not req.message:
        return JSONResponse(
            status_code=status.HTTP_400_BAD_REQUEST,
            content={"success": False, "message": "Name, email, and message are required"},
        )

    query = HelpQuery(
        id=generate_uuid(),
        name=req.name.strip(),
        email=req.email.strip().lower(),
        subject=req.subject or "No subject",
        message=req.message.strip(),
    )
    db.add(query)
    db.commit()
    db.refresh(query)

    return {
        "success": True,
        "message": "Query submitted successfully",
        "query_id": query.id,
    }
